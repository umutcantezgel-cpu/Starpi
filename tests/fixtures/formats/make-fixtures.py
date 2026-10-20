#!/usr/bin/env python3
"""Writes the Office sample files in this directory with real writers, so the extractors are tested
against the XML those programs produce and not only against XML written by the tests themselves.

    python3 tests/fixtures/formats/make-fixtures.py [--libreoffice]

Needs python-docx, openpyxl and python-pptx; --libreoffice also writes the *-lo variants by round
tripping each file through `soffice --headless --convert-to`. pandoc, when installed, writes
notes-pandoc.docx from Markdown. Document properties (author, timestamps) are reset to fixed values
so the files are reproducible and carry nothing about the machine that made them. The content is
fictitious.
"""
import datetime
import io
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

import docx
import openpyxl
from docx.shared import Pt
from openpyxl.styles import numbers
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
FIXED = datetime.datetime(2026, 1, 1, 0, 0, 0)

CORE_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    '<dc:title>Sample</dc:title>'
    '<dcterms:created xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:created>'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:modified>'
    '</cp:coreProperties>'
)
APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"></Properties>'
)


def scrub(path):
    """Rewrites docProps/core.xml and docProps/app.xml with neutral values and fixed timestamps."""
    with zipfile.ZipFile(path) as src:
        items = [(info, src.read(info.filename)) for info in src.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as out:
        for info, data in items:
            if info.filename == 'docProps/core.xml':
                data = CORE_XML.encode('utf-8')
            elif info.filename == 'docProps/app.xml':
                data = APP_XML.encode('utf-8')
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            out.writestr(fixed, data)
    with open(path, 'wb') as fh:
        fh.write(buf.getvalue())


def write_docx(path):
    d = docx.Document()
    d.add_heading('Harbor Library Renovation', level=0)
    d.add_heading('Overview', level=1)
    d.add_paragraph('The reading room reopens on 3 March 2027. The approved budget is 240,000 EUR.')
    d.add_heading('Milestones', level=2)
    d.add_paragraph('Design review in May', style='List Bullet')
    d.add_paragraph('Two bids compared', style='List Bullet 2')
    d.add_paragraph('Contract signed in June', style='List Bullet')
    d.add_paragraph('Order the shelves', style='List Number')
    d.add_paragraph('Install the lighting', style='List Number')
    d.add_heading('Team', level=2)
    table = d.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['Role', 'Name', 'Hours']):
        cell.text = text
    for row in [['Lead', 'Ana Ruiz', '120'], ['Architect', 'Ben Okafor', '80']]:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    p = d.add_paragraph('Grüße an das Team in Köln: ')
    run = p.add_run('bold part')
    run.bold = True
    run.font.size = Pt(12)
    p.add_run(' and ')
    p.add_run('a second line').add_break()
    p.add_run('after the break.')
    d.core_properties.author = ''
    d.core_properties.created = FIXED
    d.core_properties.modified = FIXED
    d.save(path)


def write_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Budget'
    ws.append(['Item', 'Cost', 'Due', 'Share'])
    ws.append(['Shelves', 42000, datetime.date(2027, 1, 15), 0.25])
    ws.append(['Lighting', 18500.5, datetime.date(2027, 2, 1), 0.125])
    ws.append(['Total', '=B2+B3', None, None])
    for row in (2, 3):
        ws.cell(row=row, column=3).number_format = 'yyyy-mm-dd'
        ws.cell(row=row, column=4).number_format = numbers.FORMAT_PERCENTAGE
    notes = wb.create_sheet('Notes')
    notes['A1'] = 'Opening'
    notes['B1'] = datetime.datetime(2027, 3, 3, 9, 30)
    notes['B1'].number_format = 'yyyy-mm-dd hh:mm'
    notes['A3'] = 'Contact'
    notes['B3'] = 'Ana Ruiz'
    notes['C3'] = True
    wb.properties.creator = ''
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.save(path)


def write_pptx(path):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = 'Harbor Library Renovation'
    s1.placeholders[1].text = 'Status update, spring 2027'
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = 'Budget'
    body = s2.placeholders[1].text_frame
    body.text = 'Approved: 240,000 EUR'
    p = body.add_paragraph()
    p.text = 'Spent so far: 61,000 EUR'
    p = body.add_paragraph()
    p.text = 'Lighting tender still open'
    p.level = 1
    s2.notes_slide.notes_text_frame.text = 'Mention that the lighting tender closes in April.'
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = 'Phases'
    shape = s3.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5))
    for r, row in enumerate([['Phase', 'Start'], ['Design', 'May 2026'], ['Build', 'September 2026']]):
        for c, text in enumerate(row):
            shape.table.cell(r, c).text = text
    prs.core_properties.author = ''
    prs.core_properties.created = FIXED
    prs.core_properties.modified = FIXED
    prs.save(path)


def libreoffice(path, ext):
    out = tempfile.mkdtemp()
    profile = 'file://' + os.path.join(out, 'profile')
    try:
        subprocess.run(['soffice', '-env:UserInstallation=' + profile, '--headless', '--convert-to', ext, '--outdir', out, path], check=True, capture_output=True)
        name = os.path.splitext(os.path.basename(path))[0] + '.' + ext
        target = os.path.join(HERE, os.path.splitext(os.path.basename(path))[0] + '-lo.' + ext)
        shutil.move(os.path.join(out, name), target)
        scrub(target)
    finally:
        shutil.rmtree(out, ignore_errors=True)


def pandoc(path):
    md = (
        '# Field Notes\n\n'
        'The survey covered *12 sites* in two weeks.\n\n'
        '## Findings\n\n'
        '- Most paths were dry\n'
        '- Two bridges need repair\n\n'
        '1. Repair the north bridge\n'
        '2. Mark the trail\n\n'
        '| Site | Visitors |\n|------|---------:|\n| North | 340 |\n| South | 125 |\n'
    )
    subprocess.run(['pandoc', '-f', 'markdown', '-o', path], input=md.encode('utf-8'), check=True)
    scrub(path)


def main():
    files = {'report.docx': write_docx, 'budget.xlsx': write_xlsx, 'status.pptx': write_pptx}
    for name, write in files.items():
        path = os.path.join(HERE, name)
        write(path)
        scrub(path)
    if '--libreoffice' in sys.argv:
        for name in files:
            libreoffice(os.path.join(HERE, name), name.rsplit('.', 1)[1])
    if shutil.which('pandoc'):
        pandoc(os.path.join(HERE, 'notes-pandoc.docx'))


if __name__ == '__main__':
    main()
